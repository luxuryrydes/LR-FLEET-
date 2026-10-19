"""GPS-vs-fuel reconciliation analytics (pandas).

Coverage analysis only: it finds gaps (km without fuel, fuel without km, missing GPS days). It deliberately
produces no mileage, cost-per-km or expense totals, which belong to the Java calculation engine.
"""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LOOKBACK_DAYS = 3


def reconcile(vehicles: pd.DataFrame, gps: pd.DataFrame, fuel: pd.DataFrame, start: date, end: date) -> dict[str, pd.DataFrame]:
    """
    vehicles: id, registration_number, status
    gps:      vehicle_id, km_date, distance_km
    fuel:     id, vehicle_id, txn_date, energy_type, has_odometer
    """
    days = (end - start).days + 1
    gps = gps.copy()
    fuel = fuel.copy()
    if not gps.empty:
        gps["km_date"] = pd.to_datetime(gps["km_date"]).dt.date
        gps["distance_km"] = pd.to_numeric(gps["distance_km"])
    if not fuel.empty:
        fuel["txn_date"] = pd.to_datetime(fuel["txn_date"]).dt.date

    g = gps.groupby("vehicle_id").agg(gps_days=("km_date", "nunique"), active_days=("distance_km", lambda s: int((s > 0).sum()))) \
        if not gps.empty else pd.DataFrame(columns=["gps_days", "active_days"])
    f = fuel.groupby("vehicle_id").agg(fuel_txns=("id", "count"), fuel_days=("txn_date", "nunique")) \
        if not fuel.empty else pd.DataFrame(columns=["fuel_txns", "fuel_days"])
    summary = vehicles.rename(columns={"id": "vehicle_id"}).set_index("vehicle_id").join(g).join(f).fillna(
        {"gps_days": 0, "active_days": 0, "fuel_txns": 0, "fuel_days": 0}).reset_index()
    for c in ("gps_days", "active_days", "fuel_txns", "fuel_days"):
        summary[c] = summary[c].astype(int)
    summary["missing_gps_days"] = days - summary["gps_days"]

    def issues(r) -> str:
        out = []
        if r.active_days > 0 and r.fuel_txns == 0:
            out.append("KM_WITHOUT_FUEL")
        if r.fuel_txns > 0 and r.gps_days == 0:
            out.append("FUEL_WITHOUT_KM")
        if r.missing_gps_days > 0 and r.status == "ACTIVE":
            out.append("MISSING_GPS_DAYS")
        return ", ".join(out)

    summary["issues"] = summary.apply(issues, axis=1) if not summary.empty else ""

    # Weeks with km but no fuel purchase
    weeks = pd.DataFrame(columns=["registration_number", "week_starting", "km_days"])
    if not gps.empty:
        active = gps[gps["distance_km"] > 0].copy()
        active["week_starting"] = active["km_date"].map(lambda d: d - timedelta(days=d.weekday()))
        wk = active.groupby(["vehicle_id", "week_starting"]).agg(km_days=("km_date", "nunique")).reset_index()
        if not fuel.empty:
            fw = fuel.copy()
            fw["week_starting"] = fw["txn_date"].map(lambda d: d - timedelta(days=d.weekday()))
            bought = set(zip(fw["vehicle_id"], fw["week_starting"]))
            wk = wk[[(v, w) not in bought for v, w in zip(wk["vehicle_id"], wk["week_starting"])]]
        weeks = wk.merge(vehicles.rename(columns={"id": "vehicle_id"})[["vehicle_id", "registration_number"]], on="vehicle_id") \
            [["registration_number", "week_starting", "km_days"]].sort_values(["registration_number", "week_starting"])

    # Fuel transactions with no GPS record on the day or the previous LOOKBACK_DAYS days and no odometer
    orphan = pd.DataFrame(columns=["transaction_id", "registration_number", "txn_date", "energy_type"])
    if not fuel.empty:
        have = set(zip(gps["vehicle_id"], gps["km_date"])) if not gps.empty else set()
        mask = []
        for v, d, odo in zip(fuel["vehicle_id"], fuel["txn_date"], fuel["has_odometer"]):
            covered = any((v, d - timedelta(days=k)) in have for k in range(LOOKBACK_DAYS + 1))
            mask.append(not covered and not bool(odo))
        o = fuel[mask].merge(vehicles.rename(columns={"id": "vehicle_id"})[["vehicle_id", "registration_number"]], on="vehicle_id")
        orphan = o.rename(columns={"id": "transaction_id"})[["transaction_id", "registration_number", "txn_date", "energy_type"]]

    summary = summary[["registration_number", "status", "gps_days", "missing_gps_days", "active_days", "fuel_txns", "fuel_days", "issues"]] \
        .sort_values("registration_number")
    return {"Summary": summary, "Km without fuel (weeks)": weeks, "Fuel without km": orphan}


def write_workbook(sheets: dict[str, pd.DataFrame], path: Path, title: str) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        notes = pd.DataFrame({"Notes": [
            title,
            "Coverage analysis only. Mileage, cost per km and expense totals come from the Java engine (Reports).",
            "KM_WITHOUT_FUEL: km recorded but no fuel purchase. FUEL_WITHOUT_KM: fuel but no GPS km.",
            f"Fuel without km: no GPS record on the transaction day or the previous {LOOKBACK_DAYS} days, and no odometer.",
        ]})
        notes.to_excel(xw, sheet_name="Read me", index=False)
        for name, df in sheets.items():
            out = df.copy()
            for c in out.columns:
                if out[c].map(lambda v: isinstance(v, date)).any():
                    out[c] = out[c].map(lambda v: v.strftime("%d-%m-%Y") if isinstance(v, date) else v)
            out.columns = [str(c).replace("_", " ").capitalize() for c in out.columns]
            out.to_excel(xw, sheet_name=name[:31], index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="DDDDDD")
                cell.alignment = Alignment(vertical="center")
            for i, col in enumerate(ws.columns, start=1):
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 80)
            ws.freeze_panes = "A2"
