"""Reading .xlsx / .xls / .csv files into plain row lists, sheet discovery and header detection."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, time
from pathlib import Path

from .mapping import FUEL_FIELDS, GPS_FIELDS, MASTER_FIELDS, header_score, suggest_mapping

MAX_HEADER_SCAN = 30
SAMPLE_ROWS = 5


class FileFormatError(Exception):
    """The file cannot be read; retrying will not help."""


def sheet_names(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext == ".csv":
        return ["CSV"]
    if ext == ".xlsx":
        from openpyxl import load_workbook
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 - surface a friendly message for any corrupt workbook
            raise FileFormatError(f"Cannot open Excel file: {e}") from e
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    if ext == ".xls":
        xlrd = _xlrd()
        return xlrd.open_workbook(str(path), on_demand=True).sheet_names()
    raise FileFormatError(f"Unsupported file type {ext}")


def read_rows(path: Path, sheet: str | None = None, limit: int | None = None) -> list[list[object]]:
    """All rows of a sheet as lists of Python values (None for empty cells). Streams .xlsx in read-only mode."""
    ext = path.suffix.lower()
    if ext == ".csv":
        return _read_csv(path, limit)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise FileFormatError(f"Cannot open Excel file: {e}") from e
        try:
            if sheet is not None and sheet not in wb.sheetnames:
                raise FileFormatError(f"Worksheet '{sheet}' not found")
            ws = wb[sheet] if sheet else wb.worksheets[0]
            rows: list[list[object]] = []
            for r in ws.iter_rows(values_only=True):
                rows.append([_clean(c) for c in r])
                if limit is not None and len(rows) >= limit:
                    break
            return _trim(rows)
        finally:
            wb.close()
    if ext == ".xls":
        xlrd = _xlrd()
        book = xlrd.open_workbook(str(path))
        sh = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
        rows = []
        for i in range(sh.nrows if limit is None else min(sh.nrows, limit)):
            row = []
            for c in sh.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(c.value, book.datemode))
                elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(_clean(c.value))
            rows.append(row)
        return _trim(rows)
    raise FileFormatError(f"Unsupported file type {ext}")


def _xlrd():
    try:
        import xlrd  # type: ignore
        return xlrd
    except ImportError as e:  # pragma: no cover
        raise FileFormatError("Legacy .xls support is not installed; please save the file as .xlsx") from e


def _read_csv(path: Path, limit: int | None) -> list[list[object]]:
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:  # pragma: no cover - latin-1 never fails
        raise FileFormatError("Cannot decode CSV file")
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    rows: list[list[object]] = []
    for r in csv.reader(io.StringIO(text), delimiter=delimiter):
        rows.append([_clean(c) for c in r])
        if limit is not None and len(rows) >= limit:
            break
    return _trim(rows)


def _clean(v: object) -> object:
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip().replace(" ", " ")
        return s if s != "" else None
    return v


def _trim(rows: list[list[object]]) -> list[list[object]]:
    """Drop fully empty trailing rows and columns."""
    while rows and all(c is None for c in rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            if r[i] is not None:
                width = max(width, i + 1)
                break
    return [list(r[:width]) + [None] * (width - len(r[:width])) for r in rows]


def detect_header_row(rows: list[list[object]]) -> int:
    """0-based index of the most header-like row among the first rows (titles above the table are skipped)."""
    best, best_score = 0, -1
    for i, r in enumerate(rows[:MAX_HEADER_SCAN]):
        s = header_score(r)
        if s > best_score:
            best, best_score = i, s
    return best


def headers_from(row: list[object]) -> list[str]:
    """Header labels; blank or duplicate labels get a positional name so every column is addressable."""
    out: list[str] = []
    seen: set[str] = set()
    for i, c in enumerate(row):
        name = " ".join(str(c).split()) if c is not None and str(c).strip() else f"Column {i + 1}"
        base, n = name, 2
        while name in seen:
            name = f"{base} ({n})"
            n += 1
        seen.add(name)
        out.append(name)
    return out


def jsonable(v: object) -> object:
    if isinstance(v, datetime):
        return v.strftime("%d-%m-%Y %H:%M:%S") if (v.hour or v.minute or v.second) else v.strftime("%d-%m-%Y")
    if isinstance(v, date):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    return str(v)


def describe_file(path: Path) -> dict:
    """Sheets with detected header row, headers, sample rows and suggested mappings (for the mapping wizard)."""
    sheets = []
    for name in sheet_names(path):
        rows = read_rows(path, None if name == "CSV" else name, limit=MAX_HEADER_SCAN + SAMPLE_ROWS + 1)
        if not rows:
            sheets.append({"name": name, "headerRow": 1, "headers": [], "sample": [], "suggestedMapping": {}})
            continue
        h = detect_header_row(rows)
        headers = headers_from(rows[h])
        sample = [dict(zip(headers, (jsonable(c) for c in r))) for r in rows[h + 1:h + 1 + SAMPLE_ROWS]]
        sheets.append({
            "name": name,
            "headerRow": h + 1,
            "headers": headers,
            "sample": sample,
            "suggestedMapping": {"GPS_KM": suggest_mapping(headers, "GPS_KM"), "FUEL": suggest_mapping(headers, "FUEL"),
                                 "VEHICLE_MASTER": suggest_mapping(headers, "VEHICLE_MASTER")},
        })
    return {"sheets": sheets, "targetFields": {"GPS_KM": GPS_FIELDS, "FUEL": FUEL_FIELDS, "VEHICLE_MASTER": MASTER_FIELDS}}
