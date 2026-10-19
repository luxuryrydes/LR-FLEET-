from datetime import date

import pandas as pd
from openpyxl import load_workbook

from fleet_worker.reconcile import reconcile, write_workbook


def test_reconciliation_finds_gaps_and_has_no_money_or_mileage(tmp_path):
    vehicles = pd.DataFrame([[1, "DL1RT1234", "ACTIVE"], [2, "MH12AB1234", "ACTIVE"], [3, "KA05MN4321", "ACTIVE"]],
                            columns=["id", "registration_number", "status"])
    gps = pd.DataFrame([[1, date(2026, 10, 1), 100], [1, date(2026, 10, 2), 120], [2, date(2026, 10, 1), 50]],
                       columns=["vehicle_id", "km_date", "distance_km"])
    fuel = pd.DataFrame([[10, 2, date(2026, 10, 2), "PETROL", False], [11, 3, date(2026, 10, 2), "CNG", False]],
                        columns=["id", "vehicle_id", "txn_date", "energy_type", "has_odometer"])
    sheets = reconcile(vehicles, gps, fuel, date(2026, 10, 1), date(2026, 10, 3))
    s = sheets["Summary"].set_index("registration_number")
    assert "KM_WITHOUT_FUEL" in s.loc["DL1RT1234", "issues"]
    assert s.loc["DL1RT1234", "missing_gps_days"] == 1
    assert "FUEL_WITHOUT_KM" in s.loc["KA05MN4321", "issues"]
    assert s.loc["MH12AB1234", "issues"].count("KM_WITHOUT_FUEL") == 0
    assert list(sheets["Fuel without km"]["transaction_id"]) == [11]
    for df in sheets.values():
        assert not any(k in " ".join(map(str, df.columns)).lower() for k in ("mileage", "amount", "cost", "expense"))

    out = tmp_path / "r.xlsx"
    write_workbook(sheets, out, "test")
    wb = load_workbook(out)
    assert wb.sheetnames[0] == "Read me"
    assert wb["Summary"]["A1"].value == "Registration number"
