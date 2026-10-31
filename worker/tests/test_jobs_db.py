"""Worker <-> PostgreSQL contract test. Runs only when WORKER_TEST_DATABASE_URL points at an empty test database
(CI provides one). Applies the backend's Flyway SQL migrations, then runs the real jobs end to end."""
import os
from pathlib import Path

import pytest

from fleet_worker.config import Config

URL = os.environ.get("WORKER_TEST_DATABASE_URL")
pytestmark = pytest.mark.skipif(not URL, reason="WORKER_TEST_DATABASE_URL not set")
MIGRATIONS = Path(__file__).resolve().parents[2] / "backend" / "src" / "main" / "resources" / "db" / "migration"


@pytest.fixture
def env(tmp_path):
    import psycopg
    from fleet_worker.jobs import connect
    from fleet_worker.storage import Storage

    with psycopg.connect(URL, autocommit=True) as c:
        c.execute("DROP SCHEMA IF EXISTS fleet CASCADE")
        c.execute("CREATE SCHEMA fleet")
        c.execute("SET search_path TO fleet")
        for f in sorted(MIGRATIONS.glob("V*.sql"), key=lambda p: int(p.name[1:].split("__")[0])):
            c.execute(f.read_text(encoding="utf-8"))
    cfg = Config(URL, "local", str(tmp_path / "files"), "", "", "", "", "", 0.1, "test", 15)
    conn = connect(cfg)
    yield conn, Storage(cfg), tmp_path
    conn.close()


def test_detect_and_parse_gps_import(env):
    from openpyxl import Workbook
    from fleet_worker.jobs import run_one

    conn, storage, tmp = env
    wb = Workbook()
    ws = wb.active
    ws.title = "Distance"
    ws.append(["Distance report"])
    ws.append(["Device Name", "Start Time", "Distance (Kms)"])
    ws.append(["DL1RT1234 (CNG)", "01-10-2026 08:00:00", 120.5])
    ws.append(["DL1RT1234 (CNG)", "01-10-2026 08:00:00", 120.5])
    ws.append(["UNKNOWN 1", "01-10-2026", 10])
    key = "import/2026/10/test-gps.xlsx"
    dest = tmp / "files" / key
    dest.parent.mkdir(parents=True)
    wb.save(dest)

    with conn.cursor() as cur:
        cur.execute("""INSERT INTO vehicles (registration_number, registration_normalized, make, model, fuel_type, category_code, ownership_type)
                       VALUES ('DL1RT1234', 'DL1RT1234', 'Maruti', 'Dzire', 'CNG', 'SEDAN', 'OWNED') RETURNING id""")
        vid = cur.fetchone()["id"]
        cur.execute("""INSERT INTO stored_files (storage_key, original_filename, content_type, size_bytes, sha256, purpose)
                       VALUES (%s, 'test-gps.xlsx', 'x', 1, repeat('0', 64), 'IMPORT') RETURNING id""", (key,))
        fid = cur.fetchone()["id"]
        cur.execute("""INSERT INTO import_batches (import_type, file_id, original_filename, file_sha256, status)
                       VALUES ('GPS_KM', %s, 'test-gps.xlsx', repeat('0', 64), 'HEADERS_PENDING') RETURNING id""", (fid,))
        bid = cur.fetchone()["id"]
        cur.execute("INSERT INTO background_jobs (job_type, payload) VALUES ('DETECT_HEADERS', %s)", (f'{{"batchId": {bid}}}',))
    conn.commit()

    assert run_one(conn, storage, "test") is True
    with conn.cursor() as cur:
        cur.execute("SELECT status, sheet_info FROM import_batches WHERE id = %s", (bid,))
        b = cur.fetchone()
    assert b["status"] == "HEADERS_READY"
    sheet = b["sheet_info"]["sheets"][0]
    assert sheet["headerRow"] == 2
    mapping = sheet["suggestedMapping"]["GPS_KM"]

    with conn.cursor() as cur:
        cur.execute("""UPDATE import_batches SET status = 'PARSE_PENDING', selected_sheet = 'Distance', header_row = 2,
                       column_mapping = %s, options = '{"recordMode": "DAILY_TOTAL"}' WHERE id = %s""",
                    (__import__("json").dumps(mapping), bid))
        cur.execute("INSERT INTO background_jobs (job_type, payload) VALUES ('PARSE_IMPORT', %s)", (f'{{"batchId": {bid}}}',))
    conn.commit()
    assert run_one(conn, storage, "test") is True

    with conn.cursor() as cur:
        cur.execute("SELECT status, total_rows, valid_rows, rejected_rows, duplicate_rows FROM import_batches WHERE id = %s", (bid,))
        b = cur.fetchone()
        cur.execute("SELECT row_number, status, vehicle_id, parsed FROM import_staging_rows WHERE batch_id = %s ORDER BY row_number", (bid,))
        rows = cur.fetchall()
        cur.execute("SELECT row_number, error_code FROM import_row_errors WHERE batch_id = %s", (bid,))
        errors = cur.fetchall()
        cur.execute("SELECT status FROM background_jobs ORDER BY id")
        jobs = [r["status"] for r in cur.fetchall()]
    assert b == {"status": "PARSED", "total_rows": 3, "valid_rows": 1, "rejected_rows": 1, "duplicate_rows": 1}
    assert [r["status"] for r in rows] == ["VALID", "DUPLICATE", "REJECTED"]
    assert rows[0]["vehicle_id"] == vid and rows[0]["parsed"]["km"] == "120.5"
    assert errors == [{"row_number": 5, "error_code": "UNMATCHED_VEHICLE"}]
    assert jobs == ["SUCCEEDED", "SUCCEEDED"]


def test_unreadable_file_fails_batch_without_retry(env):
    from fleet_worker.jobs import run_one

    conn, storage, tmp = env
    key = "import/bad.xlsx"
    (tmp / "files" / "import").mkdir(parents=True)
    (tmp / "files" / key).write_bytes(b"PK\x03\x04 not really a workbook")
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO stored_files (storage_key, original_filename, content_type, size_bytes, sha256, purpose)
                       VALUES (%s, 'bad.xlsx', 'x', 1, repeat('0', 64), 'IMPORT') RETURNING id""", (key,))
        fid = cur.fetchone()["id"]
        cur.execute("""INSERT INTO import_batches (import_type, file_id, original_filename, file_sha256, status)
                       VALUES ('FUEL', %s, 'bad.xlsx', repeat('0', 64), 'HEADERS_PENDING') RETURNING id""", (fid,))
        bid = cur.fetchone()["id"]
        cur.execute("INSERT INTO background_jobs (job_type, payload) VALUES ('DETECT_HEADERS', %s)", (f'{{"batchId": {bid}}}',))
    conn.commit()
    run_one(conn, storage, "test")
    with conn.cursor() as cur:
        cur.execute("SELECT status, error_message FROM import_batches WHERE id = %s", (bid,))
        b = cur.fetchone()
        cur.execute("SELECT status, attempts FROM background_jobs")
        j = cur.fetchone()
    assert b["status"] == "FAILED" and "Cannot open Excel file" in b["error_message"]
    assert j == {"status": "FAILED", "attempts": 1}


def test_documents_register_reads_every_tab_and_heartbeat(env):
    import json
    from datetime import datetime
    from openpyxl import Workbook
    from fleet_worker.jobs import heartbeat, run_one

    conn, storage, tmp = env
    wb = Workbook()
    ws = wb.active
    ws.title = "Commercial cars"
    ws.append(["COMMERCIAL CARS"])
    ws.append(["S. No.", "Vehicle No", "Insurance End Date", "Pullution End Date", "Permit 1 Year End Date"])
    ws.append([1, "DL52GD0247", datetime(2027, 9, 13), "NA", datetime(2026, 9, 28)])
    todo = wb.create_sheet("INSURANCE")
    todo.append(["March 2026 INSURANCE"])
    todo.append(["S.NO", "Vehical No", "INSURANCE END DATE"])
    todo.append([1, "DL52GD0247", datetime(2026, 3, 1)])
    key = "import/2026/10/test-docs.xlsx"
    dest = tmp / "files" / key
    dest.parent.mkdir(parents=True)
    wb.save(dest)
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO vehicles (registration_number, registration_normalized, make, model, fuel_type, category_code, ownership_type)
                       VALUES ('DL52GD0247', 'DL52GD0247', 'Tata', 'Tigor EV', 'ELECTRIC', 'SEDAN', 'OWNED')""")
        cur.execute("""INSERT INTO stored_files (storage_key, original_filename, content_type, size_bytes, sha256, purpose)
                       VALUES (%s, 'test-docs.xlsx', 'x', 1, repeat('0', 64), 'IMPORT') RETURNING id""", (key,))
        fid = cur.fetchone()["id"]
        cur.execute("""INSERT INTO import_batches (import_type, file_id, original_filename, file_sha256, status)
                       VALUES ('DOCUMENTS', %s, 'test-docs.xlsx', repeat('0', 64), 'HEADERS_PENDING') RETURNING id""", (fid,))
        bid = cur.fetchone()["id"]
        cur.execute("INSERT INTO background_jobs (job_type, payload) VALUES ('DETECT_HEADERS', %s)", (f'{{"batchId": {bid}}}',))
    conn.commit()
    assert run_one(conn, storage, "test") is True
    with conn.cursor() as cur:
        cur.execute("SELECT sheet_info FROM import_batches WHERE id = %s", (bid,))
        sheet = cur.fetchone()["sheet_info"]["sheets"][0]
        mapping = sheet["suggestedMapping"]["DOCUMENTS"]
        assert mapping["insuranceExpiry"] == "Insurance End Date"
        cur.execute("""UPDATE import_batches SET status = 'PARSE_PENDING', selected_sheet = 'Commercial cars', header_row = %s,
                       column_mapping = %s, options = '{"allSheets": true}' WHERE id = %s""", (sheet["headerRow"], json.dumps(mapping), bid))
        cur.execute("INSERT INTO background_jobs (job_type, payload) VALUES ('PARSE_IMPORT', %s)", (f'{{"batchId": {bid}}}',))
    conn.commit()
    assert run_one(conn, storage, "test") is True
    with conn.cursor() as cur:
        cur.execute("SELECT status, valid_rows, duplicate_rows FROM import_batches WHERE id = %s", (bid,))
        b = cur.fetchone()
        assert (b["status"], b["valid_rows"], b["duplicate_rows"]) == ("PARSED", 1, 1)   # the older to-do row is not used
        cur.execute("SELECT parsed FROM import_staging_rows WHERE batch_id = %s AND status = 'VALID'", (bid,))
        docs = {d["docType"]: d for d in cur.fetchone()["parsed"]["documents"]}
    assert docs["INSURANCE"]["expiryDate"] == "2027-09-13" and docs["PUC"]["notApplicable"]

    heartbeat(conn, "test", "host")
    with conn.cursor() as cur:
        cur.execute("SELECT count(*) AS n FROM worker_heartbeats WHERE worker_id = 'test' AND last_seen > now() - interval '1 minute'")
        assert cur.fetchone()["n"] == 1


def test_backup_and_restore_round_trip(env, tmp_path, monkeypatch):
    from fleet_worker import backup

    conn, storage, tmp = env
    monkeypatch.setenv("WORKER_DATABASE_URL", URL)
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO vehicles (registration_number, registration_normalized, make, model, fuel_type, category_code, ownership_type)
                       VALUES ('DL1BK0001', 'DL1BK0001', 'Maruti', 'Dzire', 'CNG', 'SEDAN', 'OWNED') RETURNING id""")
        vid = cur.fetchone()["id"]
        cur.execute("""INSERT INTO gps_daily_km (vehicle_id, km_date, distance_km, record_mode, source) VALUES (%s, '2026-10-01', 123.45, 'MANUAL', 'MANUAL')""", (vid,))
    conn.commit()
    files = tmp_path / "files"
    (files / "doc").mkdir(parents=True)
    (files / "doc" / "policy.pdf").write_bytes(b"%PDF-1.4 test")
    out = tmp_path / "backups"
    first = backup.create(out, keep=2, files_dir=files)
    assert first.exists()

    # change everything after the backup: a new vehicle, a deleted km row, a lost file
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO vehicles (registration_number, registration_normalized, make, model, fuel_type, category_code, ownership_type)
                       VALUES ('DL1BK0002', 'DL1BK0002', 'Maruti', 'Dzire', 'CNG', 'SEDAN', 'OWNED')""")
        cur.execute("DELETE FROM gps_daily_km")
    conn.commit()
    (files / "doc" / "policy.pdf").unlink()

    res = backup.restore(first, files)
    assert res["files"] == 1 and (files / "doc" / "policy.pdf").read_bytes() == b"%PDF-1.4 test"
    with conn.cursor() as cur:
        cur.execute("SELECT registration_normalized FROM vehicles ORDER BY id")
        assert [r["registration_normalized"] for r in cur.fetchall()] == ["DL1BK0001"]
        cur.execute("SELECT distance_km FROM gps_daily_km")
        assert str(cur.fetchone()["distance_km"]) == "123.45"
        # ids continue after the restored rows
        cur.execute("""INSERT INTO vehicles (registration_number, registration_normalized, make, model, fuel_type, category_code, ownership_type)
                       VALUES ('DL1BK0003', 'DL1BK0003', 'Maruti', 'Dzire', 'CNG', 'SEDAN', 'OWNED') RETURNING id""")
        assert cur.fetchone()["id"] > vid
    conn.commit()

    # only the newest `keep` backups stay
    import shutil
    for stamp in ("2026-01-01_0100", "2026-01-02_0100"):
        shutil.copy(first, out / f"lr-fleet-backup-{stamp}.zip")
    backup.create(out, keep=2, files_dir=None)
    assert len([p for p in out.iterdir() if p.name.startswith("lr-fleet-backup-")]) == 2
